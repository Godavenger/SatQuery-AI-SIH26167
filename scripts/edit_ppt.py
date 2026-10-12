"""Insert the two demo diagrams into the SatQuery AI idea PPT and tighten
the text so judges can read it fast.

Design rules honoured here (learned the hard way):
  * The deck is dark navy text (1B2430) on WHITE. Never set light text.
  * Never clear()+rebuild a text frame: that drops the run formatting.
    Instead edit run.text in place and delete surplus <a:p> elements, so
    size/colour/bold survive untouched.
  * Images must respect their real aspect ratio and fit the free band, or
    they silently bury the shapes underneath.

Slide 2  -> change-detection result (proves "evidence, not prose")
Slide 4  -> agent trace + class stats (proves "prototype already runs")
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC        = "SatQuery_AI_SIH26167_Idea_Submission.pptx"
IMG_CHANGE = "docs/Prototype/crops/change_core.png"          # ratio 3.16
IMG_TRACE  = "docs/Prototype/demo_sample2_agent_trace.png"   # ratio 2.69, has own title
INK        = RGBColor(0x1B, 0x24, 0x30)
MUTED      = RGBColor(0x55, 0x62, 0x72)
BORDER     = RGBColor(0xC2, 0xCE, 0xDB)

prs = Presentation(SRC)
S2, S4 = prs.slides[1], prs.slides[3]


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def set_para(p, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    runs = p.runs
    if not runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(9.6)
        r.font.color.rgb = INK
        return
    # bullet decks here put the "•  " marker in run 0 and prose in run 1
    if len(runs) >= 2 and runs[0].text.strip() in {"•", "•"} or runs[0].text.startswith("•"):
        runs[1].text = text
        for r in runs[2:]:
            r._r.getparent().remove(r._r)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


def keep_paragraphs(shape, keep):
    """Delete paragraphs beyond index `keep`, preserving those retained."""
    for p in list(shape.text_frame.paragraphs)[keep:]:
        p._p.getparent().remove(p._p)


# =====================================================================
# SLIDE 2 — trim innovation to the 3 sharpest points, add change image
# =====================================================================
inn = find(S2, "TextBox 15369")
keep_paragraphs(inn, 3)
paras = inn.text_frame.paragraphs
set_para(paras[0],
         "Cross-model conflict detection. If grounding finds no water polygon while "
         "VQA says water is present, confidence drops and the conflict is surfaced — "
         "a behaviour no single monolithic VLM can have.")
set_para(paras[1],
         "Resolution-gap engineering. Adaptation is mandated at 10 m (Sentinel); final "
         "scoring is ~0.65 m Cartosat-2S + RISAT. GSD is a conditioning input, with "
         "multi-scale augmentation across 0.3–10 m.")
set_para(paras[2],
         "Abstention as a capability. An unmet precondition yields a reasoned decline "
         "naming the failed check, instead of a confident wrong answer.")
inn.height = Inches(1.20)
print("[S2] innovation trimmed to 3 points")

# lift the reproducibility callout into the space freed
for nm, top in (("Rounded Rectangle 15370", 2.84), ("Rectangle 15371", 2.84)):
    sh = find(S2, nm)
    if sh is not None:
        sh.top = Inches(top)
        sh.height = Inches(0.50)
print("[S2] callout repositioned")

# change-detection image, true aspect 3.16
w = 4.90
img = S2.shapes.add_picture(IMG_CHANGE, Inches(7.40), Inches(3.46), width=Inches(w))
img.height = Inches(w / 3.16)
img.line.color.rgb = BORDER
img.line.width = Pt(0.75)

cap = S2.shapes.add_textbox(Inches(6.72), Inches(5.02), Inches(6.25), Inches(0.17))
cap.text_frame.word_wrap = True
r = cap.text_frame.paragraphs[0].add_run()
r.text = ("Live output — Assam: real Sentinel-2 pair, 19.5% spectral change, "
          "511 ha, 74% calibrated confidence.")
r.font.size = Pt(6.6)
r.font.color.rgb = MUTED
print("[S2] change-detection image + caption inserted")

# nudge the flow block down so it clears the table above it
for nm, dy in (("TextBox 15372", 0.06), ("Rectangle 15373", 0.06),
               ("Pentagon 15374", 0.06), ("Chevron 15375", 0.06),
               ("Chevron 15376", 0.06), ("Chevron 15377", 0.06),
               ("Chevron 15378", 0.06), ("Chevron 15379", 0.06),
               ("Chevron 15380", 0.06), ("TextBox 15381", 0.06)):
    sh = find(S2, nm)
    if sh is not None:
        sh.top = sh.top + Inches(dy)
tbl = find(S2, "Table 15366")
if tbl is not None:
    tbl.height = Inches(1.40)
print("[S2] flow block cleared from table")

# =====================================================================
# SLIDE 4 — trim feasibility to 4 points (updated to real status), add trace
# =====================================================================
fea = find(S4, "TextBox 17411")
keep_paragraphs(fea, 4)
paras = fea.text_frame.paragraphs
set_para(paras[0],
         "Prototype runs end to end today. FastAPI backend + browser GUI; all six "
         "intents execute on real Sentinel-2 and Sentinel-1 pixels across 9 sites; "
         "the router scores 5/5 on the representative queries; the PDF report renders "
         "from the trace record.")
set_para(paras[1],
         "Nothing needs inventing. Every component has published open weights or a "
         "reference implementation — the work is adaptation and orchestration.")
set_para(paras[2],
         "Fits one GPU. QLoRA 4-bit on a single card; ~50 MB adapters over one shared "
         "backbone instead of four full checkpoints.")
set_para(paras[3],
         "All data is open. BigEarthNet, VRSBench, RSVQA, CDVQA, LEVIR-CD, SECOND, "
         "DOTA, DIOR — no procurement and no licence blocker.")
fea.height = Inches(2.00)
print("[S4] feasibility trimmed to 4 points, status updated to real")

w = 4.62
img2 = S4.shapes.add_picture(IMG_TRACE, Inches(0.36), Inches(3.60), width=Inches(w))
img2.height = Inches(w / 2.69)
img2.line.color.rgb = BORDER
img2.line.width = Pt(0.75)
print("[S4] agent-trace image inserted")

prs.save(SRC)
print(f"\nsaved -> {SRC}")
