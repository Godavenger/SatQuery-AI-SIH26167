"""Rebuild slide 1 to match SIH2026-IDEA-Presentation-Format.pptx exactly.

Official template slide 1:
    Title 7     Garamond 40pt bold, schemeClr tx2, (0.36,-0.58) 11.33x2.27
    Subtitle 3  Times New Roman bold, schemeClr tx1, (1.36,0.71) 9.33x1.92
    TextBox 9   (0.36,2.27) 6.48x5.14, spAutoFit
                6 bulleted paras, Arial 24pt bold, algn=just, lnSpc 200%
                bullets: Problem Statement ID / Title / Theme /
                         PS Category / Team ID / Team Name (Registered on portal)

The current deck had replaced TextBox 9 with seven separate 10.19pt boxes plus
a "SatQuery AI" branding block. Those are removed and the template's single
bulleted TextBox is copied in verbatim (deep-copied XML, so the bullet glyph,
justification and Arial run properties are the template's own, not a guess).

Only deviation, and it is forced: the real Problem Statement Title is 120
characters, so at the template's 24pt/200% it would run off the slide. Font
drops to 16pt and line spacing to 130% -- still 1.6x larger than the deck had,
so the fields are markedly more readable than before while keeping the
template's structure, font, bullets and geometry.
"""
import copy
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

DECK = "SatQuery_AI_SIH26167_Idea_Submission.pptx"
TPL  = "SIH2026-IDEA-Presentation-Format.pptx"

# values for the six official fields, in the template's order
FIELDS = [
    "Problem Statement ID – SIH26167",
    "Problem Statement Title – SatQuery AI – An Interactive Vision-Language "
    "Assistant for Multimodal Remote Sensing Image Analysis through Text Queries",
    "Theme – Space Technology",
    "PS Category – Software",
    "Team ID – <your team ID>",
    "Team Name (Registered on portal) – SpectraX",
]

FONT_PT   = 16.0
LINE_SPC  = 1.30

deck = Presentation(DECK)
tpl  = Presentation(TPL)
S1   = deck.slides[0]
T1   = tpl.slides[0]


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


# ------------------------------------------------------------------
# 1. remove the non-template shapes the deck had added
# ------------------------------------------------------------------
drop = ["TextBox 37", "Rectangle 38", "TextBox 39", "TextBox 40", "TextBox 41",
        "TextBox 42", "TextBox 43", "TextBox 44", "TextBox 45"]
removed = []
for name in drop:
    sh = find(S1, name)
    if sh is not None:
        sh._element.getparent().remove(sh._element)
        removed.append(name)
print(f"[1] removed non-template shapes: {removed}")

# ------------------------------------------------------------------
# 2. Subtitle 3 -> template geometry + Times New Roman bold
# ------------------------------------------------------------------
sub_tpl = find(T1, "Subtitle 3")
sub     = find(S1, "Subtitle 3")
if sub is not None and sub_tpl is not None:
    sub.left, sub.top   = sub_tpl.left, sub_tpl.top
    sub.width, sub.height = sub_tpl.width, sub_tpl.height
    tf = sub.text_frame
    tf.clear()
    r = tf.paragraphs[0].add_run()
    r.text = "IDEA SUBMISSION"
    r.font.bold = True
    r.font.name = "Times New Roman"
    print(f"[2] Subtitle 3 -> template geometry "
          f"({sub.left.inches:.2f},{sub.top.inches:.2f}) "
          f"{sub.width.inches:.2f}x{sub.height.inches:.2f}, Times New Roman bold")

# ------------------------------------------------------------------
# 3. copy the template's TextBox 9 in verbatim, then fill the values
# ------------------------------------------------------------------
tb_tpl = find(T1, "TextBox 9")
spTree = S1.shapes._spTree
new_el = copy.deepcopy(tb_tpl._element)

# give it an id that cannot clash with anything already on the slide
max_id = 0
for sh in S1.shapes:
    try:
        max_id = max(max_id, int(sh._element.nvSpPr.cNvPr.get("id")))
    except Exception:
        pass
new_el.nvSpPr.cNvPr.set("id", str(max_id + 1))
new_el.nvSpPr.cNvPr.set("name", "TextBox 9")
spTree.append(new_el)

# re-resolve through the shape collection so python-pptx wraps it properly
tb = find(S1, "TextBox 9")
tb.left, tb.top     = tb_tpl.left, tb_tpl.top
tb.width, tb.height = tb_tpl.width, tb_tpl.height

paras = tb.text_frame.paragraphs
# template layout: para0 is an empty spacer, paras 1..6 carry the six bullets
bullet_paras = paras[1:1 + len(FIELDS)]
for p, value in zip(bullet_paras, FIELDS):
    runs = p.runs
    if not runs:
        continue
    runs[0].text = value
    for extra in runs[1:]:
        extra._r.getparent().remove(extra._r)
    runs[0].font.size = Pt(FONT_PT)
    runs[0].font.bold = True
    runs[0].font.name = "Arial"
    p.line_spacing = LINE_SPC
    p.alignment = PP_ALIGN.JUSTIFY

print(f"[3] TextBox 9 copied from template -> "
      f"({tb.left.inches:.2f},{tb.top.inches:.2f}) "
      f"{tb.width.inches:.2f}x{tb.height.inches:.2f}")
print(f"    {len(bullet_paras)} bullets, Arial {FONT_PT:g}pt bold, "
      f"lnSpc {LINE_SPC:g}, justified")
for p in bullet_paras:
    print(f"      - {p.runs[0].text[:78]}")

deck.save(DECK)
print(f"\nsaved -> {DECK}")
