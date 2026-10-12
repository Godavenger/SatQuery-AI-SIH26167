"""Fix slide 2 right-column overlap: the 3 innovation bullets overflow into
the reproducibility callout. Shorten them and reflow callout/image/caption
against the flow-block header at T5.26.

Vertical budget for the right column (x 6.72, w 6.25):
    1.54  bullets      -> must end by 2.90
    2.94  callout 0.50 -> 3.44
    3.52  image  1.48  -> 5.00
    5.02  caption 0.16 -> 5.18
    5.26  flow header (hard floor)
"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from pptx import Presentation
from pptx.util import Inches, Pt

SRC = "SatQuery_AI_SIH26167_Idea_Submission.pptx"
prs = Presentation(SRC)
S2 = prs.slides[1]


def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def set_para(p, text):
    runs = p.runs
    if not runs:
        return
    if runs[0].text.startswith("•") and len(runs) >= 2:
        runs[1].text = text
        for r in runs[2:]:
            r._r.getparent().remove(r._r)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)


inn = find(S2, "TextBox 15369")
paras = inn.text_frame.paragraphs
set_para(paras[0],
         "Cross-model conflict detection. Grounding finds no water but VQA says water "
         "is present -> confidence drops and the conflict is surfaced. No monolithic "
         "VLM can do this.")
set_para(paras[1],
         "Resolution-gap engineering. Trained at 10 m Sentinel, scored at ~0.65 m "
         "Cartosat-2S + RISAT. GSD is a conditioning input, 0.3-10 m augmentation.")
set_para(paras[2],
         "Abstention as a capability. An unmet precondition gives a reasoned decline "
         "naming the failed check, not a confident wrong answer.")
inn.height = Inches(1.36)
print("[S2] bullets shortened to fit 1.36 in")

for nm in ("Rounded Rectangle 15370", "Rectangle 15371"):
    sh = find(S2, nm)
    if sh is not None:
        sh.top = Inches(2.94)
        sh.height = Inches(0.50)
print("[S2] callout -> T2.94")

# the picture and caption were the last two shapes added on slide 2
pics = [sh for sh in S2.shapes if sh.shape_type == 13 and sh.left > Inches(6)]
if pics:
    pic = pics[-1]
    w = 4.68
    pic.left, pic.top = Inches(7.51), Inches(3.52)
    pic.width, pic.height = Inches(w), Inches(w / 3.16)
    print(f"[S2] image -> T3.52 {w:.2f}x{w/3.16:.2f} in")

for sh in S2.shapes:
    if sh.has_text_frame and "Live output" in sh.text_frame.text:
        sh.top, sh.height = Inches(5.02), Inches(0.16)
        print("[S2] caption -> T5.02")

prs.save(SRC)
print(f"\nsaved -> {SRC}")
